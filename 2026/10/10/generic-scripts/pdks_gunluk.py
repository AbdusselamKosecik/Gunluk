"""PDKS (ZKTeco BioTime) günlük giriş-çıkış raporu -> Excel.

Veri sunucudan WinRM ile salt okunur SELECT olarak çekilir (zkbiotime DB).
Kullanım:
    python pdks_gunluk.py --tarih 2026-10-07 [--gec 06:30] [--erken 16:30] [--out rapor.xlsx]
"""
import argparse
import datetime as dt
import json
import subprocess
import tempfile
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SERVER = "192.168.0.2"
CRED_FILE = r"$env:USERPROFILE\srv02.cred"

SQL = """
DECLARE @d date = @tarih;
WITH p AS (
  SELECT emp_code,
         COUNT(*) AS okutma,
         MIN(CASE WHEN punch_state = '0' THEN punch_time END) AS ilk_giris,
         MAX(CASE WHEN punch_state = '1' THEN punch_time END) AS son_cikis,
         MIN(punch_time) AS ilk_okutma,
         MAX(punch_time) AS son_okutma
  FROM iclock_transaction
  WHERE punch_time >= @d AND punch_time < DATEADD(day, 1, @d)
  GROUP BY emp_code
)
SELECT e.emp_code AS sicil,
       LTRIM(RTRIM(ISNULL(e.first_name, '') + ' ' + ISNULL(e.last_name, ''))) AS ad_soyad,
       ISNULL(d.dept_name, '-') AS departman,
       CONVERT(varchar(19), p.ilk_giris, 120) AS ilk_giris,
       CONVERT(varchar(19), p.son_cikis, 120) AS son_cikis,
       CONVERT(varchar(19), p.ilk_okutma, 120) AS ilk_okutma,
       CONVERT(varchar(19), p.son_okutma, 120) AS son_okutma,
       ISNULL(p.okutma, 0) AS okutma
FROM personnel_employee e
LEFT JOIN personnel_department d ON d.id = e.department_id
LEFT JOIN p ON p.emp_code = e.emp_code
WHERE e.status = 0
ORDER BY d.dept_name, ad_soyad;
"""

PS_TEMPLATE = r"""
$ErrorActionPreference = 'Stop'
$cred = Import-Clixml "{cred}"
$rows = Invoke-Command -ComputerName {server} -Credential $cred -ArgumentList $args[0], $args[1] -ScriptBlock {{
  param($sql, $tarih)
  $cn = New-Object System.Data.SqlClient.SqlConnection 'Server=localhost;Database=zkbiotime;Integrated Security=true;ApplicationIntent=ReadOnly'
  $cn.Open()
  $cmd = $cn.CreateCommand(); $cmd.CommandText = $sql; $cmd.CommandTimeout = 60
  [void]$cmd.Parameters.AddWithValue('@tarih', [datetime]$tarih)
  $t = New-Object System.Data.DataTable; $t.Load($cmd.ExecuteReader()); $cn.Close()
  $t | Select-Object sicil, ad_soyad, departman, ilk_giris, son_cikis, ilk_okutma, son_okutma, okutma
}}
$rows | Select-Object sicil, ad_soyad, departman, ilk_giris, son_cikis, ilk_okutma, son_okutma, okutma |
  ConvertTo-Json -Depth 3 | Out-File -Encoding utf8 $args[2]
"""


def fetch(tarih: str) -> list[dict]:
    with tempfile.TemporaryDirectory() as tmp:
        ps1 = Path(tmp) / "fetch.ps1"
        out = Path(tmp) / "rows.json"
        ps1.write_text(PS_TEMPLATE.format(cred=CRED_FILE, server=SERVER), encoding="utf-8-sig")
        subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(ps1), SQL, tarih, str(out)],
            check=True,
        )
        rows = json.loads(out.read_text(encoding="utf-8-sig"))
    return rows if isinstance(rows, list) else [rows]


def parse(ts):
    return dt.datetime.strptime(ts, "%Y-%m-%d %H:%M:%S") if ts else None


def durum(r, gec: dt.time, erken: dt.time) -> str:
    if r["okutma"] == 0:
        return "Gelmedi"
    giris, cikis = parse(r["ilk_giris"]), parse(r["son_cikis"])
    notlar = []
    if giris is None:
        notlar.append("Giriş yok")
    elif giris.time().replace(second=0) > gec:  # dakika bazında: 06:30:45 geç sayılmaz
        notlar.append("Geç geldi")
    if cikis is None:
        notlar.append("Çıkış yok")
    elif cikis.time().replace(second=0) < erken:
        notlar.append("Erken çıktı")
    return ", ".join(notlar) or "Normal"


FILLS = {
    "Gelmedi": "F8CBAD",
    "Normal": "C6EFCE",
}
UYARI_FILL = "FFEB9C"


def build_excel(rows, tarih, gec, erken, out: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Günlük"
    basliklar = ["Sicil", "Ad Soyad", "Departman", "İlk Giriş", "Son Çıkış", "Süre (saat)", "Okutma", "Durum"]
    ws.append(basliklar)

    ozet = defaultdict(lambda: defaultdict(int))
    for r in rows:
        giris, cikis = parse(r["ilk_giris"]), parse(r["son_cikis"])
        sure = round((cikis - giris).total_seconds() / 3600, 2) if giris and cikis and cikis > giris else None
        d = durum(r, gec, erken)
        ws.append([
            r["sicil"], r["ad_soyad"], r["departman"],
            giris.strftime("%H:%M") if giris else "",
            cikis.strftime("%H:%M") if cikis else "",
            sure, r["okutma"], d,
        ])
        fill = FILLS.get(d, UYARI_FILL)
        ws.cell(ws.max_row, 8).fill = PatternFill("solid", fgColor=fill)

        o = ozet[r["departman"]]
        o["Toplam"] += 1
        o["Geldi"] += r["okutma"] > 0
        o["Gelmedi"] += d == "Gelmedi"
        o["Geç geldi"] += "Geç geldi" in d
        o["Erken çıktı"] += "Erken çıktı" in d
        o["Eksik okutma"] += ("Giriş yok" in d) or ("Çıkış yok" in d)

    _bicimle(ws, [10, 28, 22, 10, 10, 11, 9, 24])

    ws2 = wb.create_sheet("Özet")
    kolonlar = ["Toplam", "Geldi", "Gelmedi", "Geç geldi", "Erken çıktı", "Eksik okutma"]
    ws2.append(["Departman"] + kolonlar)
    genel = defaultdict(int)
    for dep in sorted(ozet):
        ws2.append([dep] + [ozet[dep][k] for k in kolonlar])
        for k in kolonlar:
            genel[k] += ozet[dep][k]
    ws2.append(["GENEL"] + [genel[k] for k in kolonlar])
    for c in ws2[ws2.max_row]:
        c.font = Font(bold=True)
    ws2.append([])
    ws2.append([f"Tarih: {tarih} | Geç eşiği: {gec:%H:%M} | Erken çıkış eşiği: {erken:%H:%M} | Kaynak: BioTime (zkbiotime)"])
    _bicimle(ws2, [26, 9, 9, 10, 11, 12, 13])

    wb.save(out)
    return genel


def _bicimle(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--tarih", default=dt.date.today().isoformat())
    ap.add_argument("--gec", default="06:30", help="Bu saatten sonra ilk giriş = geç")
    ap.add_argument("--erken", default="16:30", help="Bu saatten önce son çıkış = erken")
    ap.add_argument("--out")
    a = ap.parse_args()

    gec = dt.datetime.strptime(a.gec, "%H:%M").time()
    erken = dt.datetime.strptime(a.erken, "%H:%M").time()
    out = Path(a.out or f"PDKS_Gunluk_{a.tarih}.xlsx")

    rows = fetch(a.tarih)
    genel = build_excel(rows, a.tarih, gec, erken, out)
    print(f"{out}  |  " + "  ".join(f"{k}: {v}" for k, v in genel.items()))


if __name__ == "__main__":
    main()
